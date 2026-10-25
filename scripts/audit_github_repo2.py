# -*- coding: utf-8 -*-
"""Second audit pass: what exactly does the pushed repo claim/hold (coverage wording, slide count, gpu_exp envs)."""
import base64
import io
import json
import urllib.request

from pptx import Presentation

API = 'https://api.github.com'
H = {'Accept': 'application/vnd.github+json', 'User-Agent': 'audit2', 'X-GitHub-Api-Version': '2022-11-28'}


def get(path):
    req = urllib.request.Request(API + path, headers=H)
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode())


def blob(sha):
    r = get(f'/repos/engulf408/FINO/git/blobs/{sha}')
    return base64.b64decode(r['content'])


tree = get('/repos/engulf408/FINO/git/trees/main?recursive=1')
blobs = {t['path']: t for t in tree['tree'] if t['type'] == 'blob'}

print('== gpu_exp paths in repo ==')
for p in sorted(blobs):
    if p.startswith('gpu_exp/'):
        print(' ', p, blobs[p]['size'], 'B')

print('\n== exp paths (CPU phase) ==')
for p in sorted(blobs):
    if p.startswith('exp/') and not p.startswith('exp_restore'):
        print(' ', p, blobs[p]['size'], 'B')

report = blob(blobs['docs/02_reproduction_report.md']['sha']).decode('utf-8', 'replace')
lines = report.splitlines()
print('\n== report stats ==')
print(' lines:', len(lines), '| chars:', len(report))
for key in ['§12', '实验覆盖', '进行中', '1 + 2', '3 个', '1/35', '3/45', '6.7%', '8.6%', '45 个任务', '35 个']:
    hits = [l.strip()[:120] for l in lines if key in l][:4]
    if hits:
        print(f'  [{key}]')
        for h in hits:
            print('   ', h)

# find coverage section
idx = next((i for i, l in enumerate(lines) if '12. 实验覆盖盘点' in l), None)
if idx is not None:
    print('\n== §12 coverage section (pushed version) ==')
    for l in lines[idx:idx + 40]:
        print(' ', l[:150])

# PPT slide count + key slide text
ppt_path = [p for p in blobs if p.startswith('ppt/') and p.endswith('.pptx')]
print('\n== PPT files ==')
for p in ppt_path:
    print(' ', p, blobs[p]['size'], 'B')
    data = blob(blobs[p]['sha'])
    try:
        prs = Presentation(io.BytesIO(data))
        print('  slides:', len(prs.slides))
        for i, s in enumerate(prs.slides, 1):
            txt = ' '.join(sh.text_frame.text.replace('\n', ' ') for sh in s.shapes
                           if sh.has_text_frame and sh.text_frame.text.strip())
            if i in (1, 18, 19) or '覆盖' in txt[:200]:
                print(f'  slide {i}:', txt[:260])
        notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
        print('  slides with notes:', notes)
    except Exception as e:  # noqa: BLE001
        print('  cannot parse:', e)
